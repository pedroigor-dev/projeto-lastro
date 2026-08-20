from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "poc" / "POC-Projeto-Lastro.docx"
ASSETS = ROOT / "docs" / "assets"
SCREENSHOT = ROOT / "docs" / "screenshots" / "projeto-lastro-dashboard.png"
SCREENSHOT_DOC = ROOT / "work" / "projeto-lastro-dashboard-doc.png"
GREEN = "65A30D"
DARK = "172016"
GRAY = "5E665F"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def set_cell_text(cell, text, bold=False, color=DARK):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    props = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    props.append(repeat)


def field(run, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, code, separate, end])


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    field(run, "PAGE")


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(document, text, first_line=True):
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(1.25)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(item)


def add_figure(document, path, caption, width=6.2):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    cap = document.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(3)
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(10)
    source = document.add_paragraph("Fonte: elaboração própria (2026).")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_after = Pt(10)
    source.runs[0].font.size = Pt(9)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], DARK)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        if widths:
            table.rows[0].cells[index].width = Cm(widths[index])
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            set_cell_text(cells[index], value)
            if widths:
                cells[index].width = Cm(widths[index])
    document.add_paragraph()
    return table


def configure(document):
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(3)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    for level, size in [(1, 15), (2, 13), (3, 12)]:
        style = document.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(GREEN if level == 1 else DARK)
        style.paragraph_format.space_before = Pt(16 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Source Code" not in [style.name for style in document.styles]:
        code = document.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(9)
        code.font.color.rgb = RGBColor.from_string(DARK)
        code.paragraph_format.left_indent = Cm(.7)
        code.paragraph_format.right_indent = Cm(.7)
        code.paragraph_format.space_before = Pt(5)
        code.paragraph_format.space_after = Pt(8)


def cover(document):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("PEDRO IGOR CAMPOS COSTA").bold = True
    document.add_paragraph("\n\n\n\n")
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PROJETO LASTRO")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(GREEN)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Prova de conceito para entrega orientada por especificações,\n"
        "evidências de qualidade e desenvolvimento agêntico"
    ).bold = True
    document.add_paragraph("\n\n\n\n\n\n")
    city = document.add_paragraph("Salvador\n2026")
    city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()


def title_page(document):
    p = document.add_paragraph("PEDRO IGOR CAMPOS COSTA")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    document.add_paragraph("\n\n")
    title = document.add_paragraph("PROJETO LASTRO")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(20)
    document.add_paragraph("\n\n")
    note = document.add_paragraph(
        "Prova de conceito apresentada como documentação técnica de portfólio, com o objetivo "
        "de demonstrar decisões arquiteturais, práticas de qualidade e integração entre Java, "
        "Spring Boot, Angular, SDD e ferramentas para desenvolvimento agêntico."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.left_indent = Cm(7)
    note.paragraph_format.line_spacing = 1.0
    note.paragraph_format.space_before = Pt(40)
    document.add_paragraph("\n\n\n")
    city = document.add_paragraph("Salvador\n2026")
    city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()


def summary_pages(document):
    add_heading(document, "RESUMO", 1)
    add_body(document,
        "Esta prova de conceito apresenta o Projeto Lastro, uma aplicação para registrar mudanças "
        "de software como especificações revisáveis e vincular critérios de aceite, evidências de "
        "qualidade e decisões de release. O sistema combina Java 21, Spring Boot 4, PostgreSQL e "
        "Angular 21. A arquitetura mantém as regras de negócio fora dos frameworks, oferece uma API "
        "REST para o fluxo operacional e publica consultas MCP somente leitura para agentes. A "
        "validação incluiu testes automatizados, regra arquitetural, análise estática, cobertura de "
        "97,22% na lógica de negócio, build de produção e execução dos três serviços em containers. "
        "O resultado não pretende substituir Jira, GitLab ou Jenkins. Seu papel é preservar o vínculo "
        "entre a intenção da mudança e a evidência que sustenta a decisão de publicar.")
    p = document.add_paragraph()
    p.add_run("Palavras-chave: ").bold = True
    p.add_run("Java. Spring Boot. Angular. Spec-Driven Development. MCP. Qualidade de software.")

    add_heading(document, "ABSTRACT", 1)
    add_body(document,
        "This proof of concept presents the Projeto Lastro platform, an application that records software changes "
        "as reviewable specifications and connects acceptance criteria, quality evidence and release "
        "decisions. The solution combines Java 21, Spring Boot 4, PostgreSQL and Angular 21. Business "
        "rules remain independent from frameworks, while REST supports the operational workflow and "
        "read-only MCP tools provide safe context to development agents. Verification covered automated "
        "tests, an architecture rule, static analysis, 97.22% business-logic coverage, a production build "
        "and container execution. Lastro complements delivery tools by preserving the link between a "
        "change intention and the evidence used to authorize its release.")
    p = document.add_paragraph()
    p.add_run("Keywords: ").bold = True
    p.add_run("Java. Spring Boot. Angular. Spec-Driven Development. MCP. Software quality.")
    document.add_page_break()

    add_heading(document, "SUMÁRIO", 1)
    entries = [
        ("1 INTRODUÇÃO", 5), ("2 PROBLEMA E OBJETIVOS", 5),
        ("3 FUNDAMENTAÇÃO TÉCNICA", 6), ("4 ARQUITETURA DA SOLUÇÃO", 6),
        ("5 IMPLEMENTAÇÃO", 7), ("6 SEGURANÇA E CONFIABILIDADE", 10),
        ("7 ESTRATÉGIA DE TESTES E QUALIDADE", 10),
        ("8 DESENVOLVIMENTO AGÊNTICO", 10),
        ("9 EXECUÇÃO DA PROVA DE CONCEITO", 11),
        ("10 DECISÕES, LIMITES E EVOLUÇÃO", 11),
        ("11 CONCLUSÃO", 12), ("REFERÊNCIAS", 12)
    ]
    for entry, page in entries:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        paragraph.add_run(f"{entry}\t{page}")
        paragraph.paragraph_format.space_after = Pt(7)
    document.add_section(WD_SECTION.NEW_PAGE)


def content(document):
    add_heading(document, "1 INTRODUÇÃO")
    add_body(document,
        "Uma entrega pode passar pelo pipeline e ainda assim chegar ao ambiente sem uma explicação "
        "clara sobre o comportamento validado. Esse problema aparece quando requisito, decisão técnica, "
        "teste e aprovação vivem em ferramentas diferentes. O Projeto Lastro foi construído para "
        "experimentar outro ponto de partida: a especificação é o contrato de engenharia, e a prontidão "
        "de release é uma conclusão derivada de evidências, não uma caixa marcada manualmente.")
    add_body(document,
        "A aplicação foi planejada a partir de uma especificação versionada antes do código. As decisões "
        "que afetam estrutura, inteligência artificial e exposição para agentes foram registradas em "
        "Architecture Decision Records (ADRs). Essa sequência permite examinar não apenas o resultado, "
        "mas também os critérios e os trade-offs adotados.")

    add_heading(document, "2 PROBLEMA E OBJETIVOS")
    add_heading(document, "2.1 Problema", 2)
    add_body(document,
        "Times de produto usam Jira para coordenação, repositórios para código, SonarQube para análise "
        "e pipelines para implantação. Essas ferramentas são úteis, mas não garantem sozinhas que cada "
        "critério da mudança tenha uma evidência identificável. A fragmentação aumenta o custo da revisão "
        "e permite que um release pareça pronto mesmo quando restam comportamentos sem validação.")
    add_heading(document, "2.2 Objetivo geral", 2)
    add_body(document,
        "Construir e validar uma aplicação full stack que preserve a rastreabilidade entre problema, "
        "solução proposta, critérios de aceite, revisão técnica, implementação e release.")
    add_heading(document, "2.3 Objetivos específicos", 2)
    add_bullets(document, [
        "modelar transições explícitas e impedir estados inválidos no domínio;",
        "apresentar bloqueadores concretos quando a entrega não está pronta;",
        "expor uma interface Angular utilizável sem duplicar regras de negócio;",
        "oferecer contexto seguro para agentes por meio de consultas MCP somente leitura;",
        "produzir evidências automatizadas de testes, cobertura, análise estática e composição;",
        "documentar riscos, decisões arquiteturais, operação e possibilidades de evolução."
    ])

    add_heading(document, "3 FUNDAMENTAÇÃO TÉCNICA")
    add_heading(document, "3.1 Spec-Driven Development", 2)
    add_body(document,
        "No desenvolvimento orientado por especificações, o comportamento esperado é escrito antes da "
        "implementação e permanece versionado com ela. No Lastro, cada critério recebe identidade e pode "
        "ser verificado de forma independente. O planner produz tarefas que citam esses identificadores, "
        "o que reduz a distância entre planejamento e teste.")
    add_heading(document, "3.2 Arquitetura hexagonal", 2)
    add_body(document,
        "O domínio não importa Spring, JPA ou classes de transporte. Serviços de aplicação coordenam casos "
        "de uso, enquanto repositório e planner são portas. Essa separação torna as regras executáveis em "
        "testes rápidos e permite trocar adaptadores sem reescrever o fluxo de negócio.")
    add_heading(document, "3.3 MCP e agentes", 2)
    add_body(document,
        "Model Context Protocol oferece uma fronteira padronizada para ferramentas consumidas por agentes. "
        "A POC publica listagem, detalhe e prontidão. Nenhuma ferramenta altera o sistema. Essa restrição "
        "reduz o impacto de um prompt incorreto e mantém decisões de fluxo nas APIs controladas.")

    add_heading(document, "4 ARQUITETURA DA SOLUÇÃO")
    add_figure(document, ASSETS / "system-context.png", "Figura 1 — Contexto do Projeto Lastro", 6.1)
    add_body(document,
        "A pessoa desenvolvedora e o revisor usam o painel Angular. A interface chama a API com uma chave "
        "de desenvolvimento; a API persiste o agregado em PostgreSQL. Pipelines registram referências de "
        "evidência, enquanto agentes consultam o mesmo estado por MCP. O Actuator publica saúde e métricas.")
    add_figure(document, ASSETS / "component-architecture.png", "Figura 2 — Componentes e dependências", 6.2)
    add_body(document,
        "As setas confirmam a direção das dependências: adaptadores conhecem portas, mas o domínio não "
        "conhece infraestrutura. Um teste com ArchUnit impede regressões nessa fronteira.")

    add_heading(document, "5 IMPLEMENTAÇÃO")
    add_heading(document, "5.1 Modelo de domínio", 2)
    add_body(document,
        "ChangeSpec é o agregado responsável pelas transições e pela avaliação de prontidão. Uma spec nasce "
        "em DRAFT, passa por revisão, aprovação e implementação. Retorno para rascunho exige motivo; "
        "aprovação exige comentário; gate aprovado exige referência; release exige todos os critérios "
        "verificados e os cinco gates aprovados.")
    add_figure(document, ASSETS / "spec-lifecycle.png", "Figura 3 — Ciclo de vida da especificação", 3.3)
    add_heading(document, "5.2 Backend", 2)
    add_body(document,
        "O backend usa Java 21 e Spring Boot 4. A API aplica Bean Validation e responde falhas de negócio "
        "com Problem Details. Flyway cria o esquema, JPA implementa a persistência e o campo de versão "
        "oferece lock otimista. A configuração desativa Open Session in View para manter o carregamento "
        "de dados dentro de transações explícitas.")
    add_heading(document, "5.3 Frontend", 2)
    add_body(document,
        "O painel Angular 21 usa componentes standalone, signals para estado derivado, Reactive Forms e "
        "HttpClient tipado. O navegador mostra as ações permitidas pelo status, mas o backend valida cada "
        "transição novamente. Essa escolha trata a interface como cliente não confiável.")
    add_figure(document, SCREENSHOT_DOC, "Figura 4 — Painel Angular executado com dados de demonstração", 6.15)

    add_heading(document, "6 SEGURANÇA E CONFIABILIDADE")
    add_body(document,
        "A POC usa chave de API comparada em tempo constante, requisições sem sessão e CORS restrito ao "
        "cliente local. O container Java executa com usuário sem privilégios. O threat model reconhece que "
        "uma chave compartilhada não basta para produção e propõe OIDC, papéis e gestão de segredos.")
    add_table(document, ["Risco", "Controle atual", "Evolução"], [
        ("Atualização concorrente", "Lock otimista", "Resposta de conflito com versão atual"),
        ("Evidência forjada", "Referência obrigatória e auditoria", "Assinatura do provedor de CI"),
        ("Agente com excesso de poder", "MCP somente leitura", "Revisão formal de novas tools"),
        ("Dependência comprometida", "Lockfile, SBOM e Dependabot", "Assinatura de imagens")
    ], [4, 5.2, 5.2])

    add_heading(document, "7 ESTRATÉGIA DE TESTES E QUALIDADE")
    add_body(document,
        "O comando Maven verify executa testes de domínio, serviços, contexto Spring e arquitetura, além de "
        "Checkstyle, SpotBugs, JaCoCo e geração de SBOM. Foram executados 17 testes sem falhas. A área de "
        "domínio e serviços atingiu 97,22% de cobertura de linhas, acima do gate de 85%. DTOs, configuração "
        "e adaptadores não entram nesse indicador; o recorte está explícito no build.")
    add_table(document, ["Evidência", "Resultado observado"], [
        ("Testes backend", "17 aprovados; nenhuma falha ou salto"),
        ("Cobertura de negócio", "97,22%; gate mínimo de 85%"),
        ("Checkstyle", "0 violações"),
        ("SpotBugs", "0 achados"),
        ("Testes Angular", "2 aprovados"),
        ("Bundle Angular", "251,68 KB bruto; 65,39 KB estimados na transferência"),
        ("Containers", "PostgreSQL, backend e frontend saudáveis")
    ], [5.3, 9.2])

    add_heading(document, "8 DESENVOLVIMENTO AGÊNTICO")
    add_body(document,
        "O repositório contém CLAUDE.md e uma skill de revisão. O acordo de trabalho exige leitura da "
        "especificação, vínculo com critérios, consulta aos ADRs e execução de evidências antes de concluir "
        "uma tarefa. Assim, a produtividade do agente é limitada por controles verificáveis, em vez de "
        "depender apenas de uma instrução extensa.")
    add_body(document,
        "O planner local é deliberadamente determinístico. Ele cria tarefas para backend, frontend, testes "
        "e documentação e referencia todos os critérios. Um modelo generativo pode ser adicionado atrás da "
        "mesma porta, desde que tenha avaliação, fallback e controle de custo.")

    add_heading(document, "9 EXECUÇÃO DA PROVA DE CONCEITO")
    add_body(document,
        "A validação seguiu quatro etapas: verificação isolada do backend; build e teste do Angular; build "
        "das imagens; inicialização do Compose com health checks. Depois, três specs de demonstração foram "
        "criadas pela própria API. A spec principal avançou até IMPLEMENTING e permaneceu bloqueada por "
        "três critérios de aceite e cinco gates ainda sem evidência.")
    add_figure(document, ASSETS / "release-sequence.png", "Figura 5 — Avaliação e confirmação de release", 6.2)
    add_body(document,
        "O cenário confirma uma decisão importante: o botão de release pode aparecer durante a implementação, "
        "mas a tentativa é recusada enquanto evaluateReadiness encontrar bloqueadores. A regra permanece "
        "atômica no servidor.")

    add_heading(document, "10 DECISÕES, LIMITES E EVOLUÇÃO")
    add_table(document, ["Decisão", "Benefício", "Custo assumido"], [
        ("Monorepo", "Uma mudança preserva código, spec e evidência", "Pipelines precisam separar caches"),
        ("Domínio sem framework", "Testes rápidos e regras explícitas", "Mapeamento JPA adicional"),
        ("Planner local", "Determinismo, privacidade e custo zero", "Menor adaptação semântica"),
        ("MCP somente leitura", "Menor superfície de risco", "Agentes não automatizam transições"),
        ("API key na POC", "Execução local simples", "Inadequada para identidades corporativas")
    ], [4.1, 5.2, 5.2])
    add_body(document,
        "As próximas evoluções recomendadas são OIDC, verificação assinada de evidências de CI, paginação, "
        "integração real com Jira, testes end-to-end e observabilidade distribuída. A adoção de um modelo de "
        "IA deve vir acompanhada de dataset de avaliação e métricas de utilidade do plano.")

    add_heading(document, "11 CONCLUSÃO")
    add_body(document,
        "O Projeto Lastro demonstrou que SDD, qualidade e agentes podem compartilhar a mesma trilha sem "
        "transformar IA em autoridade de release. A principal entrega não é um CRUD de especificações, mas "
        "um conjunto de invariantes: a mudança precisa ser revisada, cada critério precisa de verificação e "
        "cada gate precisa de evidência. Os testes, a análise estática e a execução em containers confirmam "
        "a viabilidade técnica da POC. Os limites registrados indicam com honestidade o trabalho necessário "
        "para levar o desenho a um ambiente corporativo.")

    add_heading(document, "REFERÊNCIAS")
    refs = [
        "ANGULAR. Angular documentation. Disponível em: <https://angular.dev/>. Acesso em: 20 ago. 2026.",
        "FOWLER, Martin. Patterns of Enterprise Application Architecture. Boston: Addison-Wesley, 2002.",
        "GITHUB. GitHub Actions documentation. Disponível em: <https://docs.github.com/actions>. Acesso em: 20 ago. 2026.",
        "MODEL CONTEXT PROTOCOL. Specification. Disponível em: <https://modelcontextprotocol.io/>. Acesso em: 20 ago. 2026.",
        "OWASP FOUNDATION. OWASP Application Security Verification Standard. Disponível em: <https://owasp.org/www-project-application-security-verification-standard/>. Acesso em: 20 ago. 2026.",
        "SPRING. Spring AI MCP Server Boot Starter. Disponível em: <https://docs.spring.io/spring-ai/reference/api/mcp/mcp-server-boot-starter-docs.html>. Acesso em: 20 ago. 2026.",
        "SPRING. Spring Boot reference documentation. Disponível em: <https://docs.spring.io/spring-boot/>. Acesso em: 20 ago. 2026.",
        "VERNON, Vaughn. Implementing Domain-Driven Design. Boston: Addison-Wesley, 2013."
    ]
    for ref in refs:
        paragraph = document.add_paragraph(ref)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(12)
        paragraph.paragraph_format.line_spacing = 1.0


def add_headers_and_footers(document):
    for index, section in enumerate(document.sections):
        section.different_first_page_header_footer = False
        if index == 0:
            section.header.paragraphs[0].clear()
            section.footer.paragraphs[0].clear()
            continue
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        header = section.header.paragraphs[0]
        header.text = "PROJETO LASTRO  ·  POC TÉCNICA"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.name = "Arial"
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
        page_number(section.footer.paragraphs[0])
        numbering = OxmlElement("w:pgNumType")
        numbering.set(qn("w:start"), "5")
        section._sectPr.append(numbering)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(SCREENSHOT) as screenshot:
        crop_height = min(1400, screenshot.height)
        screenshot.crop((0, 0, screenshot.width, crop_height)).save(SCREENSHOT_DOC)
    document = Document()
    configure(document)
    cover(document)
    title_page(document)
    summary_pages(document)
    content(document)
    add_headers_and_footers(document)
    properties = document.core_properties
    properties.title = "Projeto Lastro — Prova de conceito"
    properties.author = "Pedro Igor Campos Costa"
    properties.subject = "Java, Spring Boot, Angular, SDD, MCP e engenharia de software"
    properties.keywords = "Java, Spring Boot, Angular, SDD, MCP, agentes, testes"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
